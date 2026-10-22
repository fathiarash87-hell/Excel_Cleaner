import pandas as pd

df = pd.read_excel("data/raw.xlsx")

for col in ["first_name", "last_name", "city"]:
    df[col] = df[col].str.strip().str.title()

df["full_name"] = df["first_name"] + " " + df["last_name"]
before = len(df)
df = df.drop_duplicates(subset=["first_name", "last_name"]).reset_index(drop=True)
print("ردیف‌های حذف‌شده:", before - len(df))

# 1) تبدیل Salary به عدد (متن‌های غیرعددی می‌شن NaN)
df["salary"] = pd.to_numeric(df["salary"], errors="coerce")

# 2) ببینیم چند خانه خالیه
print("salary خالی:", df["salary"].isna().sum())

# 3) ردیف‌های بدون Name رو حذف کن (بدون اسم، ردیف بی‌ارزشه)
df = df.dropna(subset=["first_name", "last_name"], how="all")

# 4) Salary خالی رو با میانگین پر کن
mean_salary = df["salary"].mean()
df["salary"] = df["salary"].fillna(mean_salary).round(0)

# گزارش خلاصه: میانگین و تعداد هر شهر
summary = (
    df.groupby("city")["salary"]
    .agg(["count", "mean", "min", "max"])
    .round(0)
    .reset_index()
)
print(summary)

# ذخیره‌ی خروجی با دو شیت
with pd.ExcelWriter("clean.xlsx") as writer:
    df.to_excel(writer, sheet_name="Clean", index=False)
    summary.to_excel(writer, sheet_name="Summary", index=False)

print("فایل clean.xlsx ساخته شد")

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

def format_excel(filename):
    wb = load_workbook(filename)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = max_len + 4
        ws.freeze_panes = "A2"
    wb.save(filename)

format_excel("clean.xlsx")